#!/usr/bin/env python3
"""
make_assets.py - builds every sample file for the workshop
"Mastering Copilot in Microsoft 365" (CODED, Kuwait, 29 Sep - 1 Oct 2026).

All files are about ONE fictional company: Tamra Foods Co.
All data is invented. Output is deterministic (random.seed(2026)).

Writes every file flat into  ../../site/
Writes the planted-story answers to  ./facts.json

Run:  python3 make_assets.py
Needs: openpyxl, python-docx, python-pptx  (Python 3.9)
"""

import json
import os
import random
from collections import Counter, defaultdict
from datetime import date, timedelta

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

import docx
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from pptx import Presentation
from pptx.util import Pt as PPt

random.seed(2026)

HERE = os.path.dirname(os.path.abspath(__file__))
SITE = os.path.abspath(os.path.join(HERE, "..", "..", "site"))
FACTS_PATH = os.path.join(HERE, "facts.json")

DISCLAIMER = ("Fictional sample data for training — Tamra Foods Co. "
              "is not a real company.")
DOMAIN = "tamrafoods.example"

FACTS = {}

# FY2025 revenue by unit (company profile). Total KWD 9.6 million.
FY2025_SHARE = {"Wholesale": 50, "Caf\u00e9s": 39, "Delivery app": 11}
FY2025 = {u: 9600000 * sh // 100 for u, sh in FY2025_SHARE.items()}
assert sum(FY2025.values()) == 9600000
WRITTEN = []


def out(name):
    path = os.path.join(SITE, name)
    WRITTEN.append(name)
    return path


# ---------------------------------------------------------------------------
# Date helpers (Kuwait work week: Sunday-Thursday; weekend Friday-Saturday)
# ---------------------------------------------------------------------------
DAY_NAMES = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday",
             "Saturday", "Sunday"]


def to_workday(d):
    """Move a date forward to the next Sunday-Thursday day."""
    while d.weekday() in (4, 5):  # Friday, Saturday
        d += timedelta(days=1)
    return d


def rand_date(start, end):
    return start + timedelta(days=random.randint(0, (end - start).days))


def pct(x, nd=1):
    return round(100.0 * x, nd)


# ---------------------------------------------------------------------------
# Excel helpers
# ---------------------------------------------------------------------------
def write_table_workbook(path, sheet, table_name, headers, rows, formats,
                         about_lines, widths=None):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(headers)
    for r in rows:
        ws.append(r)
    n = len(rows) + 1
    ref = "A1:%s%d" % (get_column_letter(len(headers)), n)
    t = Table(displayName=table_name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2",
                                      showFirstColumn=False,
                                      showLastColumn=False,
                                      showRowStripes=True,
                                      showColumnStripes=False)
    ws.add_table(t)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=False)
    # number formats per column
    for ci, fmt in enumerate(formats, start=1):
        if not fmt:
            continue
        for ri in range(2, n + 1):
            ws.cell(row=ri, column=ci).number_format = fmt
    # widths
    for ci, h in enumerate(headers, start=1):
        if widths and widths.get(h):
            w = widths[h]
        else:
            longest = len(str(h))
            for r in rows:
                v = r[ci - 1]
                if v is None:
                    continue
                if isinstance(v, date):
                    s = "30 Sep 2026"
                elif isinstance(v, float):
                    s = "{:,.1f}".format(v)
                elif isinstance(v, int):
                    s = "{:,}".format(v)
                else:
                    s = str(v)
                longest = max(longest, len(s))
            w = min(max(longest + 3, 10), 60)
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.freeze_panes = "A2"

    ab = wb.create_sheet("About")
    ab["A1"] = about_lines[0]
    ab["A1"].font = Font(bold=True, size=13)
    for i, line in enumerate(about_lines[1:], start=3):
        ab.cell(row=i, column=1, value=line)
    ab.column_dimensions["A"].width = 110
    wb.active = 0
    wb.save(path)


# ---------------------------------------------------------------------------
# Word helpers
# ---------------------------------------------------------------------------
def new_doc(title, subtitle=None):
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    sec = d.sections[0]
    fp = sec.footer.paragraphs[0]
    fp.text = DISCLAIMER
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in fp.runs:
        r.font.size = Pt(8)
        r.font.italic = True
        r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    d.add_heading(title, level=0)
    if subtitle:
        p = d.add_paragraph(subtitle)
        for r in p.runs:
            r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    return d


def h1(d, t):
    d.add_heading(t, level=1)


def h2(d, t):
    d.add_heading(t, level=2)


def para(d, t, bold_lead=None):
    p = d.add_paragraph()
    if bold_lead:
        p.add_run(bold_lead).bold = True
    p.add_run(t)
    return p


def bullets(d, items, style="List Bullet"):
    for it in items:
        if isinstance(it, tuple):
            p = d.add_paragraph(style=style)
            p.add_run(it[0]).bold = True
            p.add_run(it[1])
        else:
            d.add_paragraph(it, style=style)


def table(d, headers, rows, right_cols=()):
    t = d.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        c.paragraphs[0].add_run(str(h)).bold = True
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = str(v)
            if i in right_cols:
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    d.add_paragraph()
    return t


def kwd(x):
    return "{:,.0f}".format(x)


# ===========================================================================
# 2. KPIs  (built first - the H1 review uses the same numbers)
# ===========================================================================
MONTHS = [date(2026, m, 1) for m in range(1, 9)]
UNITS = ["Cafés", "Wholesale", "Delivery app"]


def build_kpis():
    rows = []
    # Delivery app: revenue +60% Jan->Aug; margin squeezed from May.
    del_rev = [92000, 98300, 104900, 111800, 119500, 127400, 137200, 148000]
    del_margin = [0.182, 0.178, 0.176, 0.171, 0.128, 0.081, 0.027, -0.031]
    del_cos = [0.452, 0.450, 0.451, 0.449, 0.450, 0.452, 0.451, 0.450]
    # Wholesale: stable ~400k, July -21% vs June, partial recovery in August.
    whs_rev = [398600, 402300, 405100, 399400, 404200, 408500, 0, 371900]
    whs_rev[6] = round(whs_rev[5] * 0.788, -2)
    # Cafes: +2%/month.
    caf_rev = [round(310000 * (1.02 ** i) + random.uniform(-900, 900), -2)
               for i in range(8)]
    caf_sat = [4.5, 4.6, 4.6, 4.5, 4.6, 4.7, 4.6, 4.7]
    whs_sat = [4.3, 4.2, 4.3, 4.3, 4.2, 4.3, 4.1, 4.2]
    del_sat = [4.2, 4.2, 4.1, 4.2, 4.0, 3.9, 3.9, 3.8]

    for i, m in enumerate(MONTHS):
        # Cafés
        rev = int(caf_rev[i])
        cos = int(round(rev * random.uniform(0.395, 0.405), -2))
        opx = int(round(rev * random.uniform(0.375, 0.385), -2))
        orders = int(round(rev / random.uniform(3.36, 3.44)))
        rows.append([m, UNITS[0], rev, cos, opx, orders, caf_sat[i]])
        # Wholesale
        rev = int(whs_rev[i])
        cos = int(round(rev * random.uniform(0.575, 0.585), -2))
        opx = int(round(rev * random.uniform(0.215, 0.225), -2))
        orders = int(round(rev / random.uniform(305, 318)))
        rows.append([m, UNITS[1], rev, cos, opx, orders, whs_sat[i]])
        # Delivery app
        rev = int(del_rev[i])
        cos = int(round(rev * del_cos[i], -2))
        opx = int(round(rev - cos - rev * del_margin[i], -2))
        orders = int(round(rev / random.uniform(7.25, 7.45)))
        rows.append([m, UNITS[2], rev, cos, opx, orders, del_sat[i]])

    headers = ["Month", "Business unit", "Revenue (KWD)",
               "Cost of sales (KWD)", "Operating costs (KWD)", "Orders",
               "Customer satisfaction (1-5)"]
    formats = ["mmm yyyy", None, "#,##0", "#,##0", "#,##0", "#,##0", "0.0"]
    write_table_workbook(
        out("tamra-monthly-kpis.xlsx"), "KPIs", "KPIs", headers, rows, formats,
        ["Tamra Foods Co. — Monthly KPIs by business unit",
         "What this file is: monthly revenue, costs, orders and customer "
         "satisfaction for each business unit (Cafés, Wholesale, Delivery app).",
         "Period covered: January 2026 to August 2026. All amounts in Kuwaiti "
         "dinar (KWD). Customer satisfaction is an average score from 1 to 5.",
         "Operating margin = (Revenue − Cost of sales − Operating costs) "
         "÷ Revenue. No totals or formulas are included on purpose.",
         DISCLAIMER])
    return rows


def kpi_facts(rows):
    by = defaultdict(dict)
    for m, u, rev, cos, opx, orders, sat in rows:
        by[u][m] = dict(rev=rev, cos=cos, opx=opx, orders=orders, sat=sat,
                        margin=(rev - cos - opx) / rev)
    d = by["Delivery app"]
    w = by["Wholesale"]
    c = by["Cafés"]
    jan, jun, jul, aug = MONTHS[0], MONTHS[5], MONTHS[6], MONTHS[7]

    # --- checks on the planted story
    growth = d[aug]["rev"] / d[jan]["rev"] - 1
    assert 0.55 <= growth <= 0.65, growth
    assert 0.16 <= d[jan]["margin"] <= 0.20
    assert -0.05 <= d[aug]["margin"] <= -0.01
    # opex grows faster than revenue from May
    for i in range(4, 8):
        m0, m1 = MONTHS[i - 1], MONTHS[i]
        assert d[m1]["opx"] / d[m0]["opx"] > d[m1]["rev"] / d[m0]["rev"]
    first_neg = "none"
    for m in MONTHS:
        if d[m]["margin"] < 0:
            first_neg = m.strftime("%b %Y")
            break
    wchg = w[jul]["rev"] / w[jun]["rev"] - 1
    assert -0.23 <= wchg <= -0.19, wchg
    assert w[jul]["rev"] < w[aug]["rev"] < w[jun]["rev"]
    for m in MONTHS[:6]:
        assert 390000 <= w[m]["rev"] <= 415000
    cgr = [c[MONTHS[i]]["rev"] / c[MONTHS[i - 1]]["rev"] - 1 for i in range(1, 8)]
    assert all(0.012 <= g <= 0.028 for g in cgr), cgr
    avg_sat = {u: sum(v[m]["sat"] for m in MONTHS) / 8 for u, v in by.items()}
    assert max(avg_sat, key=avg_sat.get) == "Cafés"
    totals = {m.strftime("%b %Y"): sum(by[u][m]["rev"] for u in UNITS)
              for m in MONTHS}
    for v in totals.values():
        assert 790000 <= v <= 900000, totals

    h1 = {}
    for u in UNITS:
        rev = sum(by[u][m]["rev"] for m in MONTHS[:6])
        cos = sum(by[u][m]["cos"] for m in MONTHS[:6])
        opx = sum(by[u][m]["opx"] for m in MONTHS[:6])
        orders = sum(by[u][m]["orders"] for m in MONTHS[:6])
        sat = sum(by[u][m]["sat"] for m in MONTHS[:6]) / 6
        h1[u] = dict(revenue=rev, cost_of_sales=cos, operating_costs=opx,
                     operating_profit=rev - cos - opx,
                     operating_margin_pct=pct((rev - cos - opx) / rev),
                     orders=orders, avg_satisfaction=round(sat, 2))
    h1_total = {k: sum(h1[u][k] for u in UNITS)
                for k in ("revenue", "cost_of_sales", "operating_costs",
                          "operating_profit", "orders")}
    h1_total["operating_margin_pct"] = pct(h1_total["operating_profit"] /
                                           h1_total["revenue"])

    FACTS["kpis"] = {
        "file": "tamra-monthly-kpis.xlsx",
        "rows": len(rows),
        "delivery_revenue_jan": d[jan]["rev"],
        "delivery_revenue_aug": d[aug]["rev"],
        "delivery_revenue_growth_jan_to_aug_pct": pct(growth),
        "delivery_operating_margin_pct_jan": pct(d[jan]["margin"]),
        "delivery_operating_margin_pct_aug": pct(d[aug]["margin"]),
        "delivery_operating_margin_pct_by_month": {
            m.strftime("%b %Y"): pct(d[m]["margin"]) for m in MONTHS},
        "delivery_operating_costs_jan": d[jan]["opx"],
        "delivery_operating_costs_aug": d[aug]["opx"],
        "delivery_operating_costs_growth_jan_to_aug_pct":
            pct(d[aug]["opx"] / d[jan]["opx"] - 1),
        "delivery_first_negative_margin_month": first_neg,
        "wholesale_revenue_jun": w[jun]["rev"],
        "wholesale_revenue_jul": w[jul]["rev"],
        "wholesale_revenue_aug": w[aug]["rev"],
        "wholesale_jun_to_jul_change_pct": pct(wchg),
        "cafes_avg_satisfaction_jan_aug": round(avg_sat["Cafés"], 2),
        "avg_satisfaction_by_unit_jan_aug": {u: round(v, 2)
                                              for u, v in avg_sat.items()},
        "cafes_revenue_jan": c[jan]["rev"],
        "cafes_revenue_aug": c[aug]["rev"],
        "cafes_avg_monthly_growth_pct": pct(sum(cgr) / len(cgr)),
        "total_revenue_by_month": totals,
        "h1_by_unit": h1,
        "h1_total": h1_total,
    }
    return by, h1, h1_total


# ===========================================================================
# 3. Sales pipeline
# ===========================================================================
SNAP = date(2026, 9, 24)  # data "as of" date for pipeline and tickets

HOTELS = ["Pearl Gulf Hotel", "Liwan Bay Resort", "Sadu Grand Hotel",
          "Qurain Palms Hotel", "Nuwair Residence Hotel", "Failaka View Hotel",
          "Salmiya Sands Hotel", "Gulf Crescent Hotel", "Jawhara Towers Hotel",
          "Sidra Heights Hotel", "Oasis Court Hotel", "Najma Suites",
          "Blue Dhow Hotel", "Warda Garden Hotel", "Sahel Plaza Hotel",
          "Zumurrud Hotel", "Mirqab Business Hotel", "Anjafa Beach Hotel",
          "Dhow Harbour Inn", "Rawda Park Hotel", "Sharq Skyline Hotel",
          "Ghazal Suites"]
COOPS = [n + " Co-op" for n in [
    "Al-Waha", "Al-Nakheel", "Al-Rayhan", "Al-Sidr", "Al-Yasmeen",
    "Al-Murooj", "Al-Khuzama", "Al-Fanar", "Al-Mahabba", "Al-Tawfeeq",
    "Al-Ghadeer", "Al-Majd", "Al-Seef", "Al-Nahda", "Al-Rabee", "Al-Hadaf",
    "Al-Wafa", "Al-Barakah"]]
CORPS = ["Najd Tower Offices", "Gulf Horizon Holding", "Qibla Business Centre",
         "Sanad Insurance Offices", "Al-Manara Engineering",
         "Dar Al-Qalam Consulting", "Rimal Logistics", "Mersal Tech Park",
         "Bawaba Law Group", "Nukhba Advisory", "Sahab Energy Offices",
         "Najoom Media House", "Wasl Healthcare Admin",
         "Fajr Engineering Consultants", "Madar Tower Offices",
         "Ofoq Bank Operations Centre", "Tamkeen Training Institute",
         "Masar Real Estate", "Kanz Investment Offices"]
AMS = ["Fatima", "Yousef", "Noura", "Khalid", "Sara"]
REGIONS = ["Capital", "Hawalli", "Ahmadi", "Farwaniya", "Jahra",
           "Mubarak Al-Kabeer"]
LOSS_REASONS = ["Price", "Delivery terms", "Chose competitor", "No budget"]


def seg_value(seg):
    lo, hi = {"Hotel": (8000, 34000), "Co-op": (5000, 24000),
              "Corporate office": (2000, 14000)}[seg]
    return int(round(random.uniform(lo, hi) / 50.0) * 50)


def build_pipeline():
    seg_clients = {"Hotel": HOTELS[:], "Co-op": COOPS[:],
                   "Corporate office": CORPS[:]}
    for v in seg_clients.values():
        random.shuffle(v)
    counters = {k: 0 for k in seg_clients}

    def next_client(seg):
        lst = seg_clients[seg]
        c = lst[counters[seg] % len(lst)]
        counters[seg] += 1
        return c

    # Stage plan per segment.  Win rate: Hotel 17/30, Co-op 10/23, Corp 5/20
    plan = []
    plan += [("Hotel", "Won")] * 17 + [("Hotel", "Lost")] * 13
    plan += [("Co-op", "Won")] * 10 + [("Co-op", "Lost")] * 13
    plan += [("Corporate office", "Won")] * 5 + [("Corporate office", "Lost")] * 15
    open_plan = {"Hotel": ["Lead"] * 5 + ["Proposal"] * 5 + ["Negotiation"] * 5,
                 "Co-op": ["Lead"] * 5 + ["Proposal"] * 5 + ["Negotiation"] * 5,
                 "Corporate office": ["Lead"] * 5 + ["Proposal"] * 4 + ["Negotiation"] * 4}
    for seg, stages in open_plan.items():
        plan += [(seg, s) for s in stages]
    # 3 long-negotiation deals (added separately) + Darwaza Hotels (Proposal)
    assert len(plan) == 116, len(plan)

    # Lost owner plan: Khalid 13 (9 Price), others spread.
    lost_idx = [i for i, p in enumerate(plan) if p[1] == "Lost"]
    random.shuffle(lost_idx)
    lost_owner = (["Khalid"] * 13 + ["Fatima"] * 7 + ["Yousef"] * 8 +
                  ["Noura"] * 6 + ["Sara"] * 7)
    assert len(lost_owner) == len(lost_idx) == 41
    khalid_reasons = (["Price"] * 9 + ["Chose competitor"] * 2 +
                      ["Delivery terms"] + ["No budget"])
    other_reasons = (["Price"] * 7 + ["Delivery terms"] * 7 +
                     ["Chose competitor"] * 8 + ["No budget"] * 6)
    random.shuffle(other_reasons)
    lost_assign = {}
    ki = oi = 0
    for idx, am in zip(lost_idx, lost_owner):
        if am == "Khalid":
            r = khalid_reasons[ki]; ki += 1
        else:
            r = other_reasons[oi]; oi += 1
        lost_assign[idx] = (am, r)

    deals = []
    for i, (seg, stage) in enumerate(plan):
        client = next_client(seg)
        value = seg_value(seg)
        region = random.choice(REGIONS)
        reason = None
        if stage == "Lost":
            am, reason = lost_assign[i]
        else:
            am = random.choice(["Fatima", "Yousef", "Noura", "Sara",
                                "Fatima", "Yousef", "Noura", "Sara", "Khalid"])
        if stage in ("Won", "Lost"):
            created = to_workday(rand_date(date(2026, 1, 4), date(2026, 8, 13)))
            close = to_workday(created + timedelta(days=random.randint(18, 70)))
            if close >= SNAP:
                close = to_workday(SNAP - timedelta(days=random.randint(3, 20)))
                if close >= SNAP:
                    close = SNAP - timedelta(days=4)
            expected = close
            days = (SNAP - close).days
        elif stage == "Lead":
            created = to_workday(rand_date(date(2026, 7, 12), date(2026, 9, 20)))
            days = (SNAP - created).days
            expected = to_workday(created + timedelta(days=random.randint(75, 130)))
        elif stage == "Proposal":
            created = to_workday(rand_date(date(2026, 6, 1), date(2026, 9, 3)))
            entry = created + timedelta(days=random.randint(4, 18))
            if entry > SNAP:
                entry = SNAP - timedelta(days=1)
            days = min((SNAP - entry).days, 49)
            expected = to_workday(SNAP + timedelta(days=random.randint(20, 95)))
        else:  # Negotiation (normal)
            days = random.randint(4, 55)
            latest = SNAP - timedelta(days=days + 14)
            created = to_workday(rand_date(date(2026, 3, 1), latest))
            if created > latest:
                created = latest - timedelta(days=2)
                created = to_workday(created) if to_workday(created) <= latest else created - timedelta(days=2)
            expected = to_workday(SNAP + timedelta(days=random.randint(7, 70)))
        deals.append(dict(client=client, seg=seg, am=am, region=region,
                          stage=stage, value=value, created=created,
                          expected=expected, days=days, reason=reason))

    # The 3 long negotiations (the planted "stuck" deals)
    long_specs = [("Sidra Heights Hotel", "Hotel", "Yousef", "Capital", 32400, 88),
                  ("Gulf Horizon Holding", "Corporate office", "Fatima", "Capital", 29100, 74),
                  ("Blue Dhow Hotel", "Hotel", "Khalid", "Ahmadi", 24500, 67)]
    for client, seg, am, region, value, days in long_specs:
        entry = SNAP - timedelta(days=days)
        created = to_workday(entry - timedelta(days=random.randint(25, 40)))
        expected = to_workday(entry + timedelta(days=random.randint(28, 40)))
        deals.append(dict(client=client, seg=seg, am=am, region=region,
                          stage="Negotiation", value=value, created=created,
                          expected=expected, days=days, reason=None))
    # Darwaza Hotels - the prospect from the Sales account brief
    deals.append(dict(client="Darwaza Hotels", seg="Hotel", am="Noura",
                      region="Capital", stage="Proposal", value=70000,
                      created=date(2026, 8, 16), expected=date(2026, 11, 30),
                      days=(SNAP - date(2026, 9, 6)).days, reason=None))

    deals.sort(key=lambda x: (x["created"], x["client"]))

    # Make Q3 won value a clean number (adjust the last Q3 win).
    def q3_won():
        return [x for x in deals if x["stage"] == "Won"
                and date(2026, 7, 1) <= x["expected"] <= date(2026, 9, 30)]
    q3 = q3_won()
    total = sum(x["value"] for x in q3)
    target = int(round(total / 5000.0) * 5000)
    adj = sorted(q3, key=lambda x: x["value"])[-1]
    adj["value"] += target - total
    assert sum(x["value"] for x in q3_won()) == target

    rows = []
    for n, x in enumerate(deals):
        x["id"] = "D-%d" % (1001 + n)
        rows.append([x["id"], x["client"], x["seg"], x["am"], x["region"],
                     x["stage"], x["value"], x["created"], x["expected"],
                     x["days"], x["reason"]])
    headers = ["Deal ID", "Client", "Segment", "Account manager", "Region",
               "Stage", "Deal value (KWD)", "Created", "Expected close",
               "Days in stage", "Loss reason"]
    formats = [None, None, None, None, None, None, "#,##0", "d mmm yyyy",
               "d mmm yyyy", "0", None]
    write_table_workbook(
        out("tamra-sales-pipeline.xlsx"), "Pipeline", "Pipeline", headers,
        rows, formats,
        ["Tamra Foods Co. — Wholesale sales pipeline (B2B deals)",
         "What this file is: every wholesale deal with hotels, co-operative "
         "societies (co-ops) and corporate offices, with its stage and value.",
         "Period covered: deals created January to September 2026. Data as of "
         "24 Sep 2026. Deal values in Kuwaiti dinar (KWD).",
         "For Won and Lost deals, 'Expected close' is the actual close date and "
         "'Days in stage' counts days since the deal closed.",
         DISCLAIMER])
    return deals


def pipeline_facts(deals):
    assert len({d["client"] for d in deals}) == 60, len({d["client"] for d in deals})
    neg_long = [d for d in deals if d["stage"] == "Negotiation" and d["days"] > 60]
    assert len(neg_long) == 3
    tot_long = sum(d["value"] for d in neg_long)
    assert 84000 <= tot_long <= 88000
    wr = {}
    for seg in ("Hotel", "Co-op", "Corporate office"):
        won = sum(1 for d in deals if d["seg"] == seg and d["stage"] == "Won")
        lost = sum(1 for d in deals if d["seg"] == seg and d["stage"] == "Lost")
        wr[seg] = dict(won=won, lost=lost, win_rate_pct=pct(won / (won + lost)))
    assert 55 <= wr["Hotel"]["win_rate_pct"] <= 60
    assert 23 <= wr["Corporate office"]["win_rate_pct"] <= 27
    assert wr["Hotel"]["win_rate_pct"] > wr["Co-op"]["win_rate_pct"] > \
        wr["Corporate office"]["win_rate_pct"]
    lost_by_am = Counter(d["am"] for d in deals if d["stage"] == "Lost")
    assert lost_by_am.most_common(1)[0][0] == "Khalid"
    assert lost_by_am["Khalid"] > sorted(lost_by_am.values())[-2]
    kr = Counter(d["reason"] for d in deals
                 if d["stage"] == "Lost" and d["am"] == "Khalid")
    assert kr.most_common(1)[0][0] == "Price" and kr["Price"] > sum(kr.values()) / 2
    for d in deals:
        assert (d["reason"] is not None) == (d["stage"] == "Lost")
        assert date(2026, 1, 1) <= d["created"] <= date(2026, 9, 30)
        assert d["created"] + timedelta(days=d["days"]) <= SNAP or d["stage"] in ("Won", "Lost")
    q3 = [d for d in deals if d["stage"] == "Won"
          and date(2026, 7, 1) <= d["expected"] <= date(2026, 9, 30)]
    stage_counts = Counter(d["stage"] for d in deals)
    won_total = sum(d["value"] for d in deals if d["stage"] == "Won")
    open_value = sum(d["value"] for d in deals
                     if d["stage"] in ("Lead", "Proposal", "Negotiation"))
    FACTS["pipeline"] = {
        "file": "tamra-sales-pipeline.xlsx",
        "data_as_of": "24 Sep 2026",
        "number_of_deals": len(deals),
        "distinct_clients": len({d["client"] for d in deals}),
        "deals_by_stage": dict(stage_counts),
        "long_negotiation_deals_over_60_days": [
            dict(id=d["id"], client=d["client"], segment=d["seg"],
                 account_manager=d["am"], value=d["value"], days=d["days"])
            for d in sorted(neg_long, key=lambda x: -x["days"])],
        "long_negotiation_total_value": tot_long,
        "win_rate_by_segment": wr,
        "overall_win_rate_pct": pct(stage_counts["Won"] /
                                    (stage_counts["Won"] + stage_counts["Lost"])),
        "lost_deals_by_account_manager": dict(lost_by_am.most_common()),
        "khalid_loss_reasons": dict(kr.most_common()),
        "loss_reasons_all": dict(Counter(d["reason"] for d in deals
                                         if d["stage"] == "Lost").most_common()),
        "q3_2026_won_value_total": sum(d["value"] for d in q3),
        "q3_2026_won_deal_count": len(q3),
        "q3_definition": "Stage = Won and Expected close (actual close date) between 1 Jul and 30 Sep 2026",
        "total_won_value": won_total,
        "open_pipeline_value": open_value,
        "darwaza_hotels_deal": [dict(id=d["id"], stage=d["stage"], value=d["value"])
                                for d in deals if d["client"] == "Darwaza Hotels"][0],
    }


# ===========================================================================
# 4. Campaign results
# ===========================================================================
def build_campaigns():
    weeks = [date(2026, 6, 7) + timedelta(weeks=i) for i in range(13)]
    assert all(w.weekday() == 6 for w in weeks) and weeks[-1] == date(2026, 8, 30)
    # (campaign, channel): weekly spend, cost per order, avg order value
    P = {
        ("Summer Iced Coffee", "Instagram"): (300, 1.30, 3.2),
        ("Summer Iced Coffee", "TikTok"): (220, 0.85, 3.1),
        ("Summer Iced Coffee", "Snapchat"): (200, 1.45, 3.0),
        ("Summer Iced Coffee", "Google Ads"): (250, 1.60, 3.4),
        ("Summer Iced Coffee", "Email"): (60, 1.25, 5.2),
        ("Weekend Brunch", "Instagram"): (250, 2.20, 8.0),
        ("Weekend Brunch", "Snapchat"): (180, 2.50, 7.5),
        ("Weekend Brunch", "Google Ads"): (200, 2.80, 8.0),
        ("Weekend Brunch", "Email"): (50, 1.40, 8.5),
        ("Loyalty App Push", "Instagram"): (200, 1.40, 3.8),
        ("Loyalty App Push", "TikTok"): (180, 1.10, 3.5),
        ("Loyalty App Push", "Snapchat"): (150, 1.60, 3.5),
        ("Loyalty App Push", "Email"): (70, 0.90, 4.5),
        ("Back to Office", "Google Ads"): (950, 6.80, 5.8),
        ("Back to Office", "Instagram"): (300, 2.50, 6.0),
        ("Back to Office", "Snapchat"): (200, 3.00, 6.0),
        ("Back to Office", "Email"): (80, 1.60, 7.0),
    }
    CTR = {"Instagram": 0.011, "TikTok": 0.014, "Snapchat": 0.009,
           "Google Ads": 0.035, "Email": 0.12}
    CONV = {"Instagram": 0.040, "TikTok": 0.035, "Snapchat": 0.030,
            "Google Ads": 0.060, "Email": 0.080}
    active = {"Summer Iced Coffee": weeks,
              "Weekend Brunch": weeks,
              "Loyalty App Push": [w for w in weeks if w.month in (7, 8)],
              "Back to Office": [w for w in weeks if w.month == 8]}
    spike_week = date(2026, 8, 2)
    rows = []
    for w in weeks:
        for (camp, ch), (sp, cpo, aov) in P.items():
            if w not in active[camp]:
                continue
            spend = int(round(sp * random.uniform(0.88, 1.12)))
            orders = spend / (cpo * random.uniform(0.9, 1.1))
            aov_w = aov * random.uniform(0.95, 1.05)
            clicks = orders / (CONV[ch] * random.uniform(0.9, 1.1))
            impr = clicks / (CTR[ch] * random.uniform(0.9, 1.1))
            if ch == "Snapchat" and w == spike_week:
                orders *= 3.2
                clicks *= 3.0
                impr *= 2.6
            orders = int(round(orders))
            rows.append([w, camp, ch, spend, int(round(impr, -1)),
                         int(round(clicks)), orders,
                         int(round(orders * aov_w))])
    headers = ["Week starting", "Campaign", "Channel", "Spend (KWD)",
               "Impressions", "Clicks", "Orders", "Revenue (KWD)"]
    formats = ["d mmm yyyy", None, None, "#,##0", "#,##0", "#,##0", "#,##0",
               "#,##0"]
    write_table_workbook(
        out("tamra-campaign-results.xlsx"), "Campaigns", "Campaigns", headers,
        rows, formats,
        ["Tamra Foods Co. — Marketing campaign results (weekly)",
         "What this file is: weekly spend and results for each campaign and "
         "channel. Weeks start on Sunday.",
         "Period covered: weeks starting 7 Jun 2026 to 30 Aug 2026 (13 weeks). "
         "Spend and revenue in Kuwaiti dinar (KWD).",
         "For Email, 'Impressions' means emails opened.",
         DISCLAIMER])
    return rows


def campaign_facts(rows):
    agg = defaultdict(lambda: [0, 0, 0])  # spend, orders, revenue
    for w, camp, ch, sp, im, cl, od, rv in rows:
        for k in (("ch", ch), ("cc", camp, ch), ("camp", camp)):
            a = agg[k]
            a[0] += sp; a[1] += od; a[2] += rv
    chans = ["Instagram", "TikTok", "Snapchat", "Google Ads", "Email"]
    ret = {c: round(agg[("ch", c)][2] / agg[("ch", c)][0], 2) for c in chans}
    assert max(ret, key=ret.get) == "Email"
    sic_cpo = {c: round(agg[("cc", "Summer Iced Coffee", c)][0] /
                        agg[("cc", "Summer Iced Coffee", c)][1], 3) for c in chans}
    assert min(sic_cpo, key=sic_cpo.get) == "TikTok"
    combos = {k[1:]: v for k, v in agg.items() if k[0] == "cc"}
    top_combo = max(combos, key=lambda k: combos[k][0])
    assert top_combo == ("Back to Office", "Google Ads"), top_combo
    bto = combos[("Back to Office", "Google Ads")]
    assert bto[2] / bto[0] < 1
    # highest single-row spend too
    assert max(rows, key=lambda r: r[3])[1:3] == ["Back to Office", "Google Ads"]
    snap_week = defaultdict(int)
    for w, camp, ch, sp, im, cl, od, rv in rows:
        if ch == "Snapchat":
            snap_week[w] += od
    spike = snap_week[date(2026, 8, 2)]
    others = [v for k, v in snap_week.items() if k != date(2026, 8, 2)]
    avg_other = sum(others) / len(others)
    assert spike > 2 * max(others), (spike, max(others))
    FACTS["campaigns"] = {
        "file": "tamra-campaign-results.xlsx",
        "rows": len(rows),
        "weeks": "13 weeks, starting Sundays 7 Jun 2026 to 30 Aug 2026",
        "return_revenue_per_spend_by_channel": ret,
        "spend_by_channel": {c: agg[("ch", c)][0] for c in chans},
        "revenue_by_channel": {c: agg[("ch", c)][2] for c in chans},
        "orders_by_channel": {c: agg[("ch", c)][1] for c in chans},
        "summer_iced_coffee_cost_per_order_by_channel": sic_cpo,
        "back_to_office_google_ads": dict(
            spend=bto[0], revenue=bto[2], orders=bto[1],
            return_revenue_per_spend=round(bto[2] / bto[0], 2),
            note="Highest spend of any campaign x channel pair (and the highest weekly spend rows)"),
        "spend_by_campaign": {c: agg[("camp", c)][0] for c in
                              ["Summer Iced Coffee", "Weekend Brunch",
                               "Loyalty App Push", "Back to Office"]},
        "return_by_campaign": {c: round(agg[("camp", c)][2] / agg[("camp", c)][0], 2)
                               for c in ["Summer Iced Coffee", "Weekend Brunch",
                                         "Loyalty App Push", "Back to Office"]},
        "snapchat_orders_by_week": {k.strftime("%d %b %Y"): v
                                    for k, v in sorted(snap_week.items())},
        "snapchat_orders_week_2_aug": spike,
        "snapchat_orders_avg_other_weeks": round(avg_other, 1),
        "snapchat_spike_ratio": round(spike / avg_other, 2),
        "snapchat_spike_reason_not_in_data": "Influencer collaboration (not recorded in the file)",
        "total_spend": sum(r[3] for r in rows),
        "total_revenue": sum(r[7] for r in rows),
        "total_orders": sum(r[6] for r in rows),
    }


# ===========================================================================
# 5. IT tickets
# ===========================================================================
CAFE_SITES = ["Sharq", "Salmiya", "Hawally", "Jabriya", "Fahaheel", "Jahra"]
SITES = CAFE_SITES + ["HQ", "Warehouse"]
CATS = ["POS terminal", "Delivery app", "Wi-Fi / network", "Laptop / hardware",
        "Email / accounts", "Printer"]
AGENTS = ["Ahmad", "Mariam", "Omar", "Lulwa"]
BASE_HOURS = {"POS terminal": 6.0, "Delivery app": 5.0, "Wi-Fi / network": 19.0,
              "Laptop / hardware": 11.0, "Email / accounts": 2.6, "Printer": 7.0}
PRI_MULT = {"P1": 0.45, "P2": 0.75, "P3": 1.0, "P4": 1.35}
REOPEN_P = {"POS terminal": 0.07, "Delivery app": 0.07, "Wi-Fi / network": 0.10,
            "Laptop / hardware": 0.05, "Email / accounts": 0.04, "Printer": 0.27}
NOTES = {
    "POS terminal": ["Card reader not responding", "Terminal restarted, working",
                     "Receipt paper sensor fault", "Cash drawer will not open"],
    "Delivery app": ["Orders not showing on kitchen screen",
                     "Menu prices not synced", "Driver app login loop",
                     "Promo code not applied"],
    "Wi-Fi / network": ["Access point rebooted", "Guest Wi-Fi down",
                        "ISP line fault, waiting for provider",
                        "Switch replaced in back office"],
    "Laptop / hardware": ["Battery replaced", "New laptop set up",
                          "Screen flicker", "Docking station swap"],
    "Email / accounts": ["Password reset", "Account locked",
                         "Shared mailbox access added", "MFA phone changed"],
    "Printer": ["Paper jam again", "Toner replaced", "Driver reinstalled",
                "Label printer offline"],
}
FIRMWARE_NOTES = ["Terminal froze after firmware 4.2 update",
                  "Card payments failing after firmware 4.2 update",
                  "Screen freezes at checkout after firmware 4.2 update",
                  "Receipt printer stopped after firmware 4.2 update",
                  "Slow restart after firmware 4.2 update",
                  "After firmware 4.2 update terminal loses network",
                  "Rolled back? Still failing after firmware 4.2 update"]


def build_tickets():
    start, end = date(2026, 7, 1), date(2026, 9, 24)
    tickets = []

    def site_for(cat):
        if cat == "POS terminal":
            return random.choice([s for s in CAFE_SITES if s != "Fahaheel"])
        if cat == "Delivery app":
            return random.choice(CAFE_SITES + ["HQ", "HQ"])
        return random.choice(SITES)

    def pick_date(site):
        while True:
            d = rand_date(start, end)
            if site in ("HQ", "Warehouse") and d.weekday() in (4, 5):
                continue
            return d

    # Base tickets (290)
    cat_w = [("POS terminal", 22), ("Delivery app", 16), ("Wi-Fi / network", 15),
             ("Laptop / hardware", 18), ("Email / accounts", 16), ("Printer", 13)]
    cats = [c for c, w in cat_w for _ in range(w)]
    for _ in range(290):
        cat = random.choice(cats)
        site = site_for(cat)
        pr = random.choices(["P1", "P2", "P3", "P4"], [5, 20, 45, 30])[0]
        if cat == "Delivery app" and pr == "P1":
            pr = "P2"
        tickets.append(dict(opened=pick_date(site), site=site, cat=cat, pri=pr,
                            note=None))
    # Planted (a): Fahaheel POS - Jul 2, Aug 13, Sep 3
    fah = ([rand_date(date(2026, 7, 1), date(2026, 7, 31)) for _ in range(2)] +
           [rand_date(date(2026, 8, 5), date(2026, 8, 31)) for _ in range(13)] +
           [rand_date(date(2026, 9, 1), date(2026, 9, 24)) for _ in range(3)])
    fw_idx = set(random.sample(range(2, 15), 7))
    for i, d in enumerate(fah):
        note = FIRMWARE_NOTES[len([j for j in fw_idx if j < i])] if i in fw_idx else None
        tickets.append(dict(opened=d, site="Fahaheel", cat="POS terminal",
                            pri=random.choice(["P1", "P2", "P2", "P3"]) if 2 <= i < 15
                            else random.choice(["P3", "P4"]), note=note))
    # Planted (b): Delivery app P1 - 12 tickets, 8 on Thursdays
    thursdays = [d for d in (start + timedelta(days=i) for i in range((end - start).days + 1))
                 if d.weekday() == 3]
    other_days = [d for d in (start + timedelta(days=i) for i in range((end - start).days + 1))
                  if d.weekday() in (6, 0, 2, 5)]
    p1_dates = random.sample(thursdays, 8) + random.sample(other_days, 4)
    for d in p1_dates:
        # HQ works Sunday-Thursday only; cafés trade every day
        p1_sites = CAFE_SITES + (["HQ"] if d.weekday() not in (4, 5) else [])
        tickets.append(dict(opened=d, site=random.choice(p1_sites),
                            cat="Delivery app", pri="P1",
                            note=random.choice([None, "Orders stuck at 'preparing'",
                                                "App checkout down at peak"])))
    assert len(tickets) == 320

    tickets.sort(key=lambda t: (t["opened"], t["site"], t["cat"]))
    for n, t in enumerate(tickets):
        t["id"] = "T-%d" % (24001 + n)
        late = t["opened"] >= date(2026, 9, 17)
        t["status"] = "Open" if random.random() < (0.45 if late else 0.015) else "Resolved"
        if t["status"] == "Resolved":
            h = BASE_HOURS[t["cat"]] * PRI_MULT[t["pri"]] * random.uniform(0.5, 1.6)
            t["hours"] = max(0.3, round(h, 1))
            rp = REOPEN_P[t["cat"]]
            if t["note"] and "firmware" in t["note"]:
                rp = 0.2
            t["reopened"] = "Yes" if random.random() < rp else "No"
        else:
            t["hours"] = None
            t["reopened"] = "No"
        t["agent"] = random.choice(AGENTS)
        if t["note"] is None and random.random() < 0.18:
            t["note"] = random.choice(NOTES[t["cat"]])

    rows = [[t["id"], t["opened"], DAY_NAMES[t["opened"].weekday()], t["site"],
             t["cat"], t["pri"], t["status"], t["hours"], t["reopened"],
             t["agent"], t["note"]] for t in tickets]
    headers = ["Ticket ID", "Opened", "Day", "Site", "Category", "Priority",
               "Status", "Hours to resolve", "Reopened", "Assigned to", "Notes"]
    formats = [None, "d mmm yyyy", None, None, None, None, None, "0.0", None,
               None, None]
    write_table_workbook(
        out("tamra-it-tickets.xlsx"), "Tickets", "Tickets", headers, rows,
        formats,
        ["Tamra Foods Co. — IT help desk tickets",
         "What this file is: every IT ticket raised by the six cafés, head "
         "office (HQ) and the warehouse, with category, priority and outcome.",
         "Period covered: tickets opened 1 Jul 2026 to 24 Sep 2026. "
         "'Hours to resolve' is blank for tickets that are still open.",
         "Priority: P1 = critical (site cannot trade), P2 = high, P3 = normal, "
         "P4 = low.",
         DISCLAIMER],
        widths={"Notes": 52})
    return tickets


def ticket_facts(tickets):
    fah = Counter(t["opened"].strftime("%b") for t in tickets
                  if t["site"] == "Fahaheel" and t["cat"] == "POS terminal")
    pos_aug_by_site = Counter(t["site"] for t in tickets
                              if t["cat"] == "POS terminal" and t["opened"].month == 8)
    assert pos_aug_by_site.most_common(1)[0][0] == "Fahaheel"
    assert fah["Aug"] >= 3 * max(fah["Jul"], fah["Sep"])
    fw = [t for t in tickets if t["note"] and "firmware 4.2" in t["note"]]
    assert all(t["site"] == "Fahaheel" and t["cat"] == "POS terminal"
               and t["opened"].month == 8 for t in fw)
    p1d = Counter(DAY_NAMES[t["opened"].weekday()] for t in tickets
                  if t["cat"] == "Delivery app" and t["pri"] == "P1")
    assert p1d.most_common(1)[0][0] == "Thursday"
    hrs = defaultdict(list)
    for t in tickets:
        if t["hours"] is not None:
            hrs[t["cat"]].append(t["hours"])
    avg_h = {c: round(sum(v) / len(v), 1) for c, v in hrs.items()}
    assert max(avg_h, key=avg_h.get) == "Wi-Fi / network"
    reo = {}
    reo_res = {}
    for c in CATS:
        allc = [t for t in tickets if t["cat"] == c]
        res = [t for t in allc if t["status"] == "Resolved"]
        reo[c] = pct(sum(t["reopened"] == "Yes" for t in allc) / len(allc))
        reo_res[c] = pct(sum(t["reopened"] == "Yes" for t in res) / len(res))
    assert max(reo, key=reo.get) == "Printer"
    assert max(reo_res, key=reo_res.get) == "Printer"
    srt = sorted(reo.values())
    assert srt[-1] - srt[-2] >= 5
    order = ["Sunday", "Monday", "Tuesday", "Wednesday", "Thursday", "Friday",
             "Saturday"]
    fah_aug_count = fah["Aug"]
    FACTS["tickets"] = {
        "file": "tamra-it-tickets.xlsx",
        "ticket_count": len(tickets),
        "open_tickets": sum(t["status"] == "Open" for t in tickets),
        "pos_fahaheel_by_month": {"Jul": fah["Jul"], "Aug": fah["Aug"],
                                  "Sep": fah["Sep"]},
        "pos_tickets_august_by_site": dict(pos_aug_by_site.most_common()),
        "firmware_4_2_notes_count": len(fw),
        "delivery_app_p1_by_weekday": {d: p1d.get(d, 0) for d in order},
        "delivery_app_p1_total": sum(p1d.values()),
        "avg_hours_to_resolve_by_category": dict(sorted(avg_h.items(),
                                                        key=lambda x: -x[1])),
        "reopen_rate_pct_by_category_all_tickets": dict(sorted(reo.items(),
                                                               key=lambda x: -x[1])),
        "reopen_rate_pct_by_category_resolved_only": dict(sorted(reo_res.items(),
                                                                 key=lambda x: -x[1])),
        "tickets_by_category": dict(Counter(t["cat"] for t in tickets).most_common()),
    }
    return fah_aug_count


# ===========================================================================
# 1. Company profile  (NO profit figures)
# ===========================================================================
def build_company_profile():
    d = new_doc("Tamra Foods Co. — Company Profile",
                "Premium dates and specialty coffee from Kuwait, since 2009")
    h1(d, "About us")
    para(d, "Tamra Foods Co. is a Kuwaiti food company. We started in 2009 "
            "with one small shop and a love for good dates. Today our head "
            "office and warehouse are in the Shuwaikh Industrial Area. We "
            "employ about 420 people.")
    para(d, "We bring together two things Kuwait loves: dates and coffee. We "
            "serve guests in our own cafés, supply hotels and shops across "
            "the country, and deliver to homes and offices through our app.")
    h1(d, "What we sell")
    bullets(d, [("Premium dates", " — Khalas, Sukkari, Ajwa and Medjool, packed in Kuwait."),
                ("Date paste", " — for bakeries, hotel kitchens and home cooks."),
                ("Stuffed dates", " — with almonds, pistachio, orange peel and tahini."),
                ("Date syrup", " — pure syrup in 250 ml and 1 litre bottles."),
                ("Specialty coffee", " — single-origin beans, roasted every week.")])
    para(d, "Coming in October 2026: Tamra Cold Brew, a cold-brew coffee "
            "sweetened with our own date syrup.")
    h1(d, "Our story")
    bullets(d, ["2009 — Tamra opens a small date shop in Kuwait City.",
                "2013 — The first Tamra Café opens in Sharq.",
                "2016 — We start supplying hotels and co-ops (Wholesale).",
                "2019 — We move to a larger warehouse and head office in Shuwaikh.",
                "2024 — We launch the Tamra delivery app.",
                "2026 — Six cafés, about 420 employees, and a new product: Tamra Cold Brew."])
    h1(d, "How we work")
    para(d, "We buy dates directly from farms we know in the Gulf region. Every "
            "batch is checked, sorted and packed in our Shuwaikh facility. We "
            "roast coffee in small batches every week, so it is always fresh. "
            "Our teams speak Arabic and English, and we serve customers seven "
            "days a week in our cafés and in the app.")
    h1(d, "Our three business units")
    h2(d, "Cafés")
    para(d, "We run six Tamra Cafés: Sharq (Kuwait City), Salmiya, Hawally, "
            "Jabriya, Fahaheel and Jahra. Each café serves coffee, date "
            "desserts and light breakfast, and sells our packed products.")
    h2(d, "Wholesale")
    para(d, "Our wholesale team supplies more than 40 hotels, 25 co-operative "
            "societies (co-ops) and 60 corporate offices. We deliver from our "
            "Shuwaikh warehouse, usually the next working day.")
    h2(d, "Delivery app")
    para(d, "We launched the Tamra app in 2024. It now has about 38,000 active "
            "users. Customers order coffee, dates and gift boxes for delivery "
            "to their home or office.")
    h1(d, "Our customers")
    bullets(d, [("Hotels", " — welcome dates for rooms, breakfast buffets and events."),
                ("Co-ops", " — packed dates and syrup on the shelves of local co-operative societies."),
                ("Corporate offices", " — coffee and date boxes for meetings, guests and staff."),
                ("Café guests", " — families, students and professionals who visit our six cafés."),
                ("App users", " — people who order to their home or office through the Tamra app.")])
    h1(d, "Revenue by business unit (FY2025)")
    para(d, "Total revenue in FY2025 was KWD 9.6 million.")
    table(d, ["Business unit", "Revenue (KWD)", "Share"],
          [[u, kwd(v), "%d%%" % FY2025_SHARE[u]] for u, v in FY2025.items()] +
          [["Total", kwd(sum(FY2025.values())), "100%"]], right_cols=(1, 2))
    FACTS["fy2025_revenue_by_unit"] = dict(
        FY2025, total=sum(FY2025.values()), share_pct=FY2025_SHARE)
    h1(d, "Leadership")
    table(d, ["Name", "Role"],
          [["Dana Al-Kandari", "Chief Executive Officer (CEO)"],
           ["Faisal Al-Rashed", "Chief Financial Officer (CFO)"],
           ["Yousef Al-Shammari", "Chief Operating Officer (COO)"],
           ["Noura Al-Ajmi", "Head of Sales"],
           ["Sara Al-Hajri", "Head of Marketing"],
           ["Omar Al-Enezi", "Head of IT"]])
    h1(d, "Our values")
    bullets(d, [("Quality", " — we choose the best dates and beans, and we never cut corners."),
                ("Hospitality", " — we treat every customer like a guest in our home."),
                ("Speed", " — we deliver on time and answer every question quickly.")])
    h1(d, "Contact")
    para(d, "Tamra Foods Co., Shuwaikh Industrial Area, Kuwait")
    para(d, "Email: info@%s  |  Web: www.%s" % (DOMAIN, DOMAIN))
    d.save(out("tamra-company-profile.docx"))


# ===========================================================================
# 6. H1 2026 business review (Executive) - numbers from the KPI rows
# ===========================================================================
def build_h1_review(by, h1r, h1t):
    C, W, D = "Cafés", "Wholesale", "Delivery app"
    jan, may, jun = MONTHS[0], MONTHS[4], MONTHS[5]
    d = new_doc("H1 2026 Business Review",
                "Tamra Foods Co. — prepared for the Board of Directors, "
                "September 2026. Period: January to June 2026.")
    h1(d, "Headline results")
    whs_top5 = round(h1r[W]["revenue"] * 0.31, -3)
    bullets(d, [
        "Revenue for H1 2026 was KWD %s." % kwd(h1t["revenue"]),
        "Operating profit was KWD %s, an operating margin of %.1f%%."
        % (kwd(h1t["operating_profit"]), h1t["operating_margin_pct"]),
        "Cafés grew every month: from KWD %s in January to KWD %s in June."
        % (kwd(by[C][jan]["rev"]), kwd(by[C][jun]["rev"])),
        "Delivery app revenue grew %.0f%% from January to June, but its margin fell "
        "from %.1f%% to %.1f%%."
        % (pct(by[D][jun]["rev"] / by[D][jan]["rev"] - 1, 0),
           pct(by[D][jan]["margin"]), pct(by[D][jun]["margin"])),
        "Wholesale stayed stable at about KWD 400,000 per month.",
        "Café guests are the most satisfied customers (average score %.1f out of 5)."
        % h1r[C]["avg_satisfaction"],
    ])
    para(d, "At this pace, full-year revenue will pass the FY2025 total of "
            "KWD 9.6 million.")
    h1(d, "Results by business unit")
    rows = []
    for u in (W, C, D):
        r = h1r[u]
        rows.append([u, kwd(r["revenue"]), kwd(r["cost_of_sales"]),
                     kwd(r["operating_costs"]), kwd(r["operating_profit"]),
                     "%.1f%%" % r["operating_margin_pct"],
                     "%.1f" % r["avg_satisfaction"]])
    rows.append(["Total", kwd(h1t["revenue"]), kwd(h1t["cost_of_sales"]),
                 kwd(h1t["operating_costs"]), kwd(h1t["operating_profit"]),
                 "%.1f%%" % h1t["operating_margin_pct"], ""])
    table(d, ["Business unit", "Revenue (KWD)", "Cost of sales (KWD)",
              "Operating costs (KWD)", "Operating profit (KWD)", "Margin",
              "Satisfaction"], rows, right_cols=(1, 2, 3, 4, 5, 6))
    para(d, "All amounts are for January to June 2026 and match the monthly "
            "KPI file (tamra-monthly-kpis.xlsx).")
    h1(d, "What is driving the numbers")
    h2(d, "1. Delivery app: growth with a margin squeeze")
    para(d, "More people use the app every month. Revenue rose from KWD %s in "
            "January to KWD %s in June. But from May, operating costs grew "
            "faster than revenue. Driver pay went up and the delivery platform "
            "raised its fees. The operating margin fell from %.1f%% in April to "
            "%.1f%% in May and %.1f%% in June."
         % (kwd(by[D][jan]["rev"]), kwd(by[D][jun]["rev"]),
            pct(by[D][MONTHS[3]]["margin"]), pct(by[D][may]["margin"]),
            pct(by[D][jun]["margin"])))
    para(d, "If this trend continues, the delivery app will lose money on every "
            "order before the end of the year.")
    h2(d, "2. Wholesale: stable, but concentrated")
    para(d, "Wholesale is our largest unit. Revenue is steady, but a small "
            "number of clients bring a large share. Our top 5 hotel clients "
            "brought about 31%% of wholesale revenue in H1 (about KWD %s)."
         % kwd(whs_top5))
    table(d, ["Top 5 hotel clients", "Share of wholesale revenue"],
          [["Pearl Gulf Hotel", "8.2%"], ["Liwan Bay Resort", "6.9%"],
           ["Sadu Grand Hotel", "6.1%"], ["Jawhara Towers Hotel", "5.3%"],
           ["Qurain Palms Hotel", "4.5%"], ["Top 5 total", "31.0%"]],
          right_cols=(1,))
    para(d, "Losing even one of these clients would hurt wholesale results.")
    h2(d, "3. Cafés: steady growth")
    para(d, "Café revenue grew about 2% every month. Guests like the new "
            "breakfast menu and the date desserts. Satisfaction stayed between "
            "4.5 and 4.7.")
    h1(d, "Outlook for H2 2026")
    bullets(d, [
        ("Tamra Cold Brew launch. ", "Our new cold brew, sweetened with date syrup, launches on 18 October 2026."),
        ("POS upgrade. ", "We plan to replace the point-of-sale (POS) terminals in all six cafés and connect them to the delivery app. Pilot at Salmiya, then rollout from October to December 2026."),
        ("Delivery-partner fees. ", "We are renegotiating fees with our delivery partner and testing a new delivery fee model for customers."),
        ("Wholesale. ", "Hotels are quieter in summer. We expect a dip in July and a recovery from September."),
    ])
    h1(d, "Risks")
    bullets(d, [
        "Delivery margin keeps falling if fees and driver pay are not fixed.",
        "Wholesale depends on a few large hotel clients.",
        "Old POS terminals fail more often and slow down service at peak times.",
        "The cold brew launch competes with other new summer and winter drinks in the market.",
        "Coffee bean prices may rise in H2.",
    ])
    h1(d, "Decisions needed from the Board")
    table(d, ["#", "Decision", "Amount / detail"],
          [["1", "Approve the POS upgrade for all six cafés", "KWD 120,000"],
           ["2", "Approve a new delivery fee model (customer delivery fee plus a lower partner commission)", "Details in the CFO paper"],
           ["3", "Approve the Tamra Cold Brew launch budget", "KWD 18,000"]])
    para(d, "Prepared by the Office of the CEO. Questions: ceo.office@%s" % DOMAIN)
    d.save(out("tamra-h1-2026-business-review.docx"))


# ===========================================================================
# 7. Account brief - Darwaza Hotels (Sales)
# ===========================================================================
def build_account_brief():
    d = new_doc("Account Brief: Darwaza Hotels",
                "Tamra Foods Co. — Wholesale team. Prepared for Noura "
                "Al-Ajmi, Head of Sales. September 2026.")
    h1(d, "About the client")
    para(d, "Darwaza Hotels is a Kuwaiti hotel group with four hotels. Together "
            "they have about 780 rooms and host many weddings and business events.")
    table(d, ["Hotel", "Area", "Rooms", "Notes"],
          [["Darwaza Grand", "Sharq, Kuwait City", "310", "Flagship; large ballroom"],
           ["Darwaza Marina Salmiya", "Salmiya", "210", "Busy weekend brunch"],
           ["Darwaza Suites Jabriya", "Jabriya", "140", "Long-stay guests, families"],
           ["Darwaza Airport Inn", "Farwaniya", "120", "Short stays, early breakfast"]],
          right_cols=(2,))
    h1(d, "Their needs")
    bullets(d, [
        ("Welcome-amenity dates. ", "A small box of premium dates in every room, every day."),
        ("Breakfast buffet. ", "Date paste, date syrup, stuffed dates and fresh dates for the buffet in all four hotels."),
        ("Event catering. ", "Date platters and coffee for weddings, conferences and Ramadan events."),
        ("Consistent quality. ", "The same look and taste in every hotel."),
    ])
    h1(d, "Problems with their current supplier")
    bullets(d, [
        "Late deliveries — about one in five deliveries arrived late in the last six months.",
        "Inconsistent quality — dates in the same order were different sizes and colours.",
        "No single contact person — each hotel calls a different number.",
        "Slow to replace damaged boxes.",
    ])
    h1(d, "Decision makers")
    table(d, ["Name", "Role", "What they care about"],
          [["Tareq Al-Mutairi", "Group Director of Operations (final decision)", "Reliability, one supplier for all hotels"],
           ["Hessa Al-Otaibi", "Group Executive Chef", "Taste, quality, product range"],
           ["Bader Al-Fadhli", "Procurement Manager", "Price, payment terms, contract"],
           ["Reem Al-Qattan", "Guest Experience Manager", "Presentation of welcome amenities"]])
    h1(d, "Budget and timeline")
    bullets(d, [
        "Budget: about KWD 70,000 per year for all four hotels.",
        "Proposal due: 15 October 2026.",
        "Tasting session: early November 2026 at Darwaza Grand.",
        "Decision: by 30 November 2026.",
        "Contract start: January 2027 (first delivery Sunday 3 January 2027).",
    ])
    h1(d, "Our offer")
    h2(d, "Products")
    bullets(d, ["Premium dates for welcome amenities (Sukkari and Medjool)",
                "Buffet range: date paste, date syrup, stuffed dates, fresh Khalas",
                "Event platters and Tamra specialty coffee"])
    h2(d, "Three price tiers")
    table(d, ["Tier", "What is included", "Price per year (KWD)"],
          [["Classic", "Room amenities + buffet range", "58,000"],
           ["Premium", "Classic + event platters (up to 40 events)", "69,500"],
           ["Signature", "Premium + custom Darwaza-branded boxes + coffee for events", "79,000"]],
          right_cols=(2,))
    h2(d, "Delivery promise")
    bullets(d, ["Next-day delivery to all four hotels.",
                "98% on-time delivery over the last 12 months.",
                "One account manager and one phone line for the whole group."])
    h1(d, "Proof points")
    bullets(d, [
        ("Pearl Gulf Hotel ", "— client for 3 years. Welcome dates in all rooms. Renewed in 2026."),
        ("Liwan Bay Resort ", "— moved to Tamra in 2025 after late deliveries from another supplier. No missed delivery since."),
        ("Al-Waha Co-op ", "— Tamra dates are a top-5 product on their shelves."),
    ])
    h1(d, "Risks and objections to prepare for")
    table(d, ["Objection or risk", "How we answer"],
          [["“Your price is higher than our current supplier.”", "Show the cost of late deliveries and waste. Offer the Classic tier."],
           ["“Can you supply four hotels at once?”", "Explain our Shuwaikh warehouse capacity and next-day routes."],
           ["“We need 60-day payment terms.”", "Our standard for hotels is 30 days. Offer 2% discount for payment within 10 days."],
           ["Their current supplier may cut prices to keep them.", "Focus on quality and reliability, not only price."],
           ["The decision may slip past 30 November.", "Agree the tasting date now and confirm the decision date in writing."]])
    para(d, "Account lead: Noura Al-Ajmi, Head of Sales (sales@%s)" % DOMAIN)
    d.save(out("tamra-account-brief-darwaza-hotels.docx"))


# ===========================================================================
# 8. Campaign brief - Tamra Cold Brew (Marketing)
# ===========================================================================
def build_cold_brew_brief():
    d = new_doc("Campaign Brief: Tamra Cold Brew Launch",
                "Tamra Foods Co. — Marketing team. Owner: Sara Al-Hajri, "
                "Head of Marketing. September 2026.")
    h1(d, "The product")
    para(d, "Tamra Cold Brew is cold-brew coffee sweetened with our own date "
            "syrup. It comes in a 250 ml bottle. It is sold in all six Tamra "
            "Cafés, in the Tamra app, and in selected co-ops.")
    h1(d, "Objective")
    para(d, "In the first 6 weeks after launch (18 October to 28 November 2026):")
    bullets(d, ["Sell 25,000 bottles.",
                "Win 4,000 new app users.",
                "Make cold brew 15% of all café coffee orders.",
                "Reach 2 million impressions on social media."])
    h1(d, "Audience")
    para(d, "Young professionals in Kuwait, aged 22 to 35. They work in offices, "
            "order coffee on their phone, and follow local food accounts. They "
            "want good coffee that feels local and modern.")
    h1(d, "Key messages")
    bullets(d, [("Real coffee, naturally sweet. ", "Slow cold-brewed coffee, sweetened with date syrup."),
                ("Made in Kuwait. ", "From our Shuwaikh kitchen to your desk."),
                ("Ready when you are. ", "Grab it in our cafés or order it in the app.")],
            style="List Number")
    h1(d, "Channels")
    bullets(d, ["Instagram — launch posts, reels and stories",
                "TikTok — short videos with local creators",
                "Snapchat — filters and ads around office areas",
                "Influencers — 5 local food and lifestyle creators",
                "In-café — free tasting, posters and table cards",
                "Email and app push — to our 38,000 app users"])
    h1(d, "Budget: KWD 18,000")
    table(d, ["Item", "KWD", "Share"],
          [["Instagram ads", "4,500", "25%"],
           ["TikTok ads and creators", "4,000", "22%"],
           ["Influencer collaborations", "3,000", "17%"],
           ["Snapchat ads", "2,500", "14%"],
           ["In-café sampling and materials", "2,500", "14%"],
           ["Contingency", "1,000", "6%"],
           ["Email and app push", "500", "3%"],
           ["Total", "18,000", "100%"]], right_cols=(1, 2))
    h1(d, "Timeline")
    table(d, ["Dates", "Phase", "What happens"],
          [["4–17 Oct", "Teaser", "“Something new is brewing” posts; app users join a waiting list"],
           ["18 Oct", "Launch day", "Launch in all six cafés and the app; creators post"],
           ["18–31 Oct", "Weeks 1–2", "Free tasting in cafés; paid social at full weight"],
           ["1–14 Nov", "Weeks 3–4", "Office push: Snapchat around office areas; bundle offer"],
           ["15–28 Nov", "Weeks 5–6", "Loyalty offer in the app; results review on 30 Nov"]])
    h1(d, "KPIs")
    table(d, ["KPI", "Target (6 weeks)"],
          [["Bottles sold", "25,000"], ["New app users", "4,000"],
           ["Cold brew share of café coffee orders", "15%"],
           ["Social impressions", "2,000,000"],
           ["Return on ad spend (revenue ÷ spend)", "3.0 or higher"]],
          right_cols=(1,))
    h1(d, "Creative ideas")
    bullets(d, ["“My desk, my Tamra” — creators show their morning routine with a bottle of cold brew at work.",
                "Street tasting — a Tamra cart gives free samples near office towers in Sharq and Shuwaikh.",
                "App-only launch bundle — two bottles and a date bar at a launch price for the first two weeks.",
                "In-café “first sip” photo wall with the Tamra logo in English and Arabic."])
    h1(d, "Risks")
    bullets(d, ["Supply: if demand is high, bottles may sell out in the first week. Production plans a 20% buffer.",
                "Price: KWD 1.750 per bottle may feel high. The launch bundle helps people try it.",
                "Weather: cold drinks sell less when it gets cooler in late November.",
                "Budget: if the budget is cut, we lose reach with the 22–35 audience on TikTok."])
    h1(d, "Brand rules")
    bullets(d, ["Bilingual: every post and poster in English and Arabic.",
                "No health claims. Do not say “healthy”, “sugar-free” or “good for you”.",
                "Tone: local and warm. Talk like a friend, not like an advert.",
                "Show real Tamra cafés and real Kuwaiti places.",
                "Use the Tamra logo and colours from the brand kit."])
    para(d, "Contact: marketing@%s" % DOMAIN)
    d.save(out("tamra-campaign-brief-cold-brew.docx"))


# ===========================================================================
# 9. POS upgrade project plan (Technical)
# ===========================================================================
def build_pos_plan(fah_aug):
    d = new_doc("Project Plan: POS Upgrade for Tamra Cafés",
                "Tamra Foods Co. — IT department. Project owner: Omar "
                "Al-Enezi, Head of IT. Version 1.0, September 2026.")
    h1(d, "Why we need this project")
    para(d, "Our point-of-sale (POS) terminals are more than six years old. "
            "They fail more often, and they do not connect to the delivery app. "
            "Staff must type app orders into the POS by hand.")
    bullets(d, [
        "POS terminal faults are one of the largest groups of IT tickets from the cafés.",
        "In August 2026, Fahaheel logged %d POS tickets. Many started after the "
        "firmware 4.2 update: terminals froze and card payments failed." % fah_aug,
        "App orders typed by hand cause mistakes and slow service at peak times.",
        "The vendor will stop supporting the old terminals at the end of 2027.",
    ])
    h1(d, "Scope")
    h2(d, "In scope")
    bullets(d, ["New POS terminals and card readers in all six cafés "
                "(Sharq, Salmiya, Hawally, Jabriya, Fahaheel, Jahra).",
                "Kitchen display screens in every café.",
                "Integration with the Tamra delivery app, so app orders go "
                "straight to the POS and the kitchen screen.",
                "Staff training and a support plan."])
    h2(d, "Out of scope")
    bullets(d, ["Warehouse and head office systems.",
                "Changes to the delivery app for customers."])
    h1(d, "Phases and timeline")
    table(d, ["Phase", "Dates (2026)", "What happens"],
          [["0. Preparation", "4–22 Oct", "Final vendor contract, network checks, order hardware"],
           ["1. Pilot at Salmiya", "25 Oct–12 Nov", "Install in Salmiya; run old and new side by side for one week"],
           ["2. Pilot review", "15–19 Nov", "Check results; fix issues; go / no-go decision"],
           ["3. Rollout wave 1", "22 Nov–3 Dec", "Fahaheel, Sharq, Hawally"],
           ["4. Rollout wave 2", "6–17 Dec", "Jabriya, Jahra"],
           ["5. Close", "20–31 Dec", "Remove old terminals; final report"]])
    para(d, "We install at night, after the café closes, so trading is not "
            "affected. Fahaheel is first in wave 1 because it had the most faults.")
    h1(d, "Budget: KWD 120,000")
    table(d, ["Item", "Detail", "KWD"],
          [["POS terminals and card readers", "36 units × KWD 1,650", "59,400"],
           ["Kitchen display screens", "12 units × KWD 900", "10,800"],
           ["Software licences (year 1)", "36 licences", "14,400"],
           ["Delivery app integration", "Vendor development and testing", "16,000"],
           ["Installation and network work", "6 cafés", "7,200"],
           ["Training", "Trainer time and materials", "4,200"],
           ["Contingency", "About 7%", "8,000"],
           ["Total", "", "120,000"]], right_cols=(2,))
    h1(d, "Risks and mitigations")
    table(d, ["Risk", "Impact", "Mitigation"],
          [["New terminals fail during a busy shift", "High", "Keep old terminals as backup for 2 weeks in each café"],
           ["Firmware problems like the August 4.2 update", "High", "Test every update at Salmiya first; no automatic updates"],
           ["Delivery app integration is late", "Medium", "Pilot can start without integration; add it before wave 1"],
           ["Staff are not confident with the new system", "Medium", "Train every shift; one “POS champion” per café"],
           ["Hardware arrives late", "Medium", "Order in week 1; vendor has stock in Kuwait"],
           ["Budget overrun", "Low", "KWD 8,000 contingency; monthly cost check with Finance"]])
    h1(d, "Support and training plan")
    bullets(d, ["Two-hour hands-on training for every café team before go-live.",
                "One POS champion per café as the first contact for questions.",
                "IT support on site for the first three days in each café.",
                "Help desk: P1 POS tickets answered within 30 minutes, 7 days a week.",
                "Short how-to videos in English and Arabic."])
    h1(d, "Success measures")
    table(d, ["Measure", "Today", "Target by March 2027"],
          [["POS tickets per café per month", "About 5", "2 or fewer"],
           ["App orders typed by hand", "All", "None"],
           ["Average checkout time", "95 seconds", "60 seconds"],
           ["Order errors on app orders", "About 3%", "Under 1%"],
           ["Staff who feel confident with the POS", "—", "90% or more"]])
    para(d, "Contact: it.support@%s" % DOMAIN)
    d.save(out("tamra-pos-upgrade-project-plan.docx"))


# ===========================================================================
# 10. Messy deck
# ===========================================================================
def build_messy_deck():
    prs = Presentation()
    L_TITLE, L_CONTENT = prs.slide_layouts[0], prs.slide_layouts[1]

    def content(title, items):
        s = prs.slides.add_slide(L_CONTENT)
        s.shapes.title.text = title
        tf = s.placeholders[1].text_frame
        tf.text = items[0]
        for it in items[1:]:
            tf.add_paragraph().text = it
        return s

    s = prs.slides.add_slide(L_TITLE)
    s.shapes.title.text = "Café Operations Review — Q3 2026"
    s.placeholders[1].text = "Tamra Foods Co. | Operations team | September 2026"
    s.notes_slide.notes_text_frame.text = DISCLAIMER

    content("Summary", [
        "Wait times are the main problem in Q3 and the main reason for complaints.",
        "Fahaheel had the worst month in August because of POS problems.",
        "Better peak-hour staffing and the new POS will fix most issues.",
        "We ask leadership to approve the Salmiya POS pilot and the new rota."])

    content("Customer complaints", [
        "In Q3 we received 1,240 customer complaints across the six cafés, which is 18% more than in Q2, and most of them came through app reviews and Instagram messages, although some also came in person at the counter.",
        "The biggest group of complaints was about waiting too long for drinks at peak times, especially between 7:30 and 9:00 in the morning and again after 8:00 in the evening on Thursdays.",
        "Many guests said that app orders were not ready when the driver arrived, so drivers waited in the café and the food and coffee were cold when they reached the customer.",
        "Some guests complained that the card machine did not work and they had to pay in cash or wait, and this happened most often at Fahaheel in August.",
        "A smaller number of complaints were about the temperature of iced drinks in summer, because ice melted quickly while guests waited at the pickup counter.",
        "Guests at Jahra and Jabriya said that parking was hard to find on weekends, which is not something we control but it still affects how they feel about the visit.",
        "Around 90 complaints were about missing items in app orders, mostly desserts and extra date syrup, which we think happens when staff type app orders into the POS by hand.",
        "We also received complaints about the new breakfast menu prices, but these were fewer after the first two weeks of July.",
        "Positive comments were mostly about staff friendliness, the date desserts and the cleanliness of the cafés, which shows the problem is speed and not hospitality."])

    content("Agenda", ["Customer complaints", "Findings: wait times",
                       "Recommendations", "Next steps", "Summary"])

    wait = [
        "The average wait time from order to drink across all cafés in Q3 was 7.8 minutes, compared with our target of 5 minutes, and it was worse on Thursdays and at weekends.",
        "Morning peak (7:30–9:00) is the worst period, with average waits above 10 minutes at Sharq and Salmiya where many office workers come before work.",
        "App orders and walk-in orders compete for the same baristas, and at peak times app orders can be 40% of all drinks being made.",
        "Staff rotas are the same every day of the week, so we have too few baristas on Thursday evenings and weekends and too many on quiet afternoons.",
        "When POS terminals freeze, the queue stops completely and staff take orders on paper, which adds about 3 minutes to each order.",
        "Cafés with the newest espresso machines (Hawally and Jabriya) have shorter waits, by about 1.5 minutes on average.",
        "Guests who wait more than 10 minutes give a satisfaction score about 0.6 lower than guests who wait less than 5 minutes."]
    content("Findings: wait times", wait)
    content("Findings: wait times (v2)", wait + [
        "Fahaheel had the longest waits in August (average 11.5 minutes), when POS terminals froze after the firmware 4.2 update."])

    content("PLACEHOLDER — delete me", [
        "Lorem ipsum dolor sit amet", "Add chart here??",
        "old numbers from Q1 — do not use", "TBC"])

    content("Recommendations", [
        "Change the staff rota so that we have one more barista on Thursday evenings and on Friday and Saturday mornings in every café, and one fewer barista on quiet weekday afternoons, which keeps total hours about the same.",
        "Create a separate app-order station with its own barista at Sharq and Salmiya during the morning peak, so that app orders do not slow down the walk-in queue.",
        "Approve the POS upgrade, starting with a pilot at Salmiya, so that app orders go straight to the kitchen screen and staff stop typing them by hand.",
        "Stop all automatic POS firmware updates until IT has tested them at one café first, so that the August problem at Fahaheel does not happen again.",
        "Replace the two oldest espresso machines at Fahaheel and Jahra in Q1 2027.",
        "Send guests who waited more than 10 minutes a free-drink voucher in the app to protect satisfaction scores."])

    content("Next steps", [
        "Yousef: new rota ready by 15 Oct",
        "Omar: Salmiya POS pilot starts 25 Oct",
        "Operations: app-order station trial at Sharq from 1 Nov",
        "Review wait times again in the Q4 review"])

    prs.save(out("tamra-messy-deck.pptx"))
    FACTS["messy_deck"] = {
        "file": "tamra-messy-deck.pptx",
        "original_order": ["1 Title", "2 Summary", "3 Customer complaints",
                           "4 Agenda", "5 Findings: wait times",
                           "6 Findings: wait times (v2)",
                           "7 PLACEHOLDER — delete me", "8 Recommendations",
                           "9 Next steps"],
        "intended_order": ["Title (slide 1)", "Agenda (slide 4)",
                           "Customer complaints (slide 3, split or trim the text wall)",
                           "Findings: wait times (merge slides 5 and 6; keep v2's extra Fahaheel point)",
                           "Recommendations (slide 8, trim the text wall)",
                           "Next steps (slide 9)", "Summary (slide 2, moved to the end)"],
        "delete": ["Slide 7 PLACEHOLDER — delete me",
                   "Slide 5 or 6 (duplicate; keep one merged slide)"],
        "intended_slide_count": 7,
        "faults": ["Summary at the start instead of the end",
                   "Agenda is slide 4 instead of slide 2",
                   "Text walls on slides 3, 5, 6 and 8",
                   "Slides 5 and 6 are near-duplicates (v2 adds one point)",
                   "Slide 7 is a junk placeholder"],
    }


# ===========================================================================
# 11. House-style letter
# ===========================================================================
def build_house_letter():
    d = new_doc("Tamra Customer Care — House-Style Letter",
                "Tone sample: use this letter as the model for Tamra's voice.")
    para(d, "Dear [Customer name],")
    para(d, "Thank you for writing to us. We are glad you told us, and we are "
            "sorry that your experience was not what you expected from Tamra.")
    para(d, "We take this seriously. It was our job to get it right, and this "
            "time we did not. There is no excuse, and we will not make one.")
    para(d, "Here is what we are doing:")
    bullets(d, ["We have shared your message with the team involved.",
                "We are checking what went wrong so it does not happen again.",
                "We have added a thank-you credit to your Tamra app account."])
    para(d, "We would like to hear from you directly. Please call our Customer "
            "Care team on 1800-TAMRA (1800-82672), Sunday to Thursday, 8:00 am "
            "to 8:00 pm. You can also email care@%s. Please mention your "
            "reference number so we can help you quickly." % DOMAIN)
    para(d, "You are the reason we do what we do. We hope to welcome you back "
            "soon for a coffee and dates on us.")
    para(d, "With warm regards,")
    p = para(d, "The Tamra Customer Care Team")
    p.runs[0].bold = True
    para(d, "Tamra Foods Co. | 1800-TAMRA | care@%s | www.%s" % (DOMAIN, DOMAIN))
    h2(d, "Our voice in one line")
    para(d, "Warm, clear and short. We say sorry once, we own the problem, we "
            "say what we are doing, and we give a direct way to reach us.")
    d.save(out("tamra-house-style-letter.docx"))


# ===========================================================================
# 12. Supplier notice (~500 words)
# ===========================================================================
def build_supplier_notice():
    d = new_doc("Internal Notice: Changes to Wholesale Delivery Cut-off Times "
                "and Payment Terms",
                "To: Sales and Customer Service staff | From: Yousef "
                "Al-Shammari, COO, and Faisal Al-Rashed, CFO | Effective: "
                "Sunday 1 November 2026")
    h2(d, "1. Purpose")
    para(d, "This notice explains two changes for wholesale customers: new order "
            "cut-off times and new payment terms. Both start on 1 November 2026. "
            "Please read it and use it when you talk to customers.")
    para(d, "Why are we doing this? Order volumes grew strongly this year. Late "
            "orders put pressure on the warehouse, and some deliveries left "
            "late. At the same time, some co-ops pay very late, which affects "
            "our cash flow. These changes protect our 98% on-time record and "
            "keep our prices stable.")
    h2(d, "2. New order cut-off times")
    para(d, "Today, orders placed by 5:00 pm arrive the next working day. From 1 "
            "November, the cut-off for next-day delivery moves to 2:00 pm. This "
            "gives the warehouse more time to pick and check orders.")
    bullets(d, ["Orders by 2:00 pm (Sunday to Thursday): delivered the next working day.",
                "Orders after 2:00 pm: delivered in two working days.",
                "Orders on Friday or Saturday: processed on Sunday and delivered on Monday."])
    para(d, "Urgent same-day orders for hotels are still possible for events, "
            "with manager approval and a delivery fee.")
    h2(d, "3. Delivery windows by region")
    table(d, ["Region", "Delivery window", "Days"],
          [["Capital", "7:00 am – 11:00 am", "Sunday – Thursday"],
           ["Hawalli", "7:00 am – 11:00 am", "Sunday – Thursday"],
           ["Farwaniya", "9:00 am – 1:00 pm", "Sunday – Thursday"],
           ["Mubarak Al-Kabeer", "9:00 am – 1:00 pm", "Sunday – Thursday"],
           ["Ahmadi", "10:00 am – 2:00 pm", "Sunday, Tuesday, Thursday"],
           ["Jahra", "10:00 am – 2:00 pm", "Monday, Wednesday"]])
    h2(d, "4. New payment terms for co-ops and hotels")
    bullets(d, ["Co-ops: payment within 45 days of the invoice (today it is 60 days).",
                "Hotels: payment stays at 30 days.",
                "Early-payment discount: 2% for any invoice paid within 10 days.",
                "Invoices unpaid after 15 days past the due date: new orders "
                "are paused until the account is up to date."])
    para(d, "Corporate office terms do not change.")
    h2(d, "5. What this means for customers")
    para(d, "Most customers will not notice a big change if they order before "
            "2:00 pm. Customers in Ahmadi and Jahra must plan orders around "
            "their delivery days. Co-ops will pay 15 days sooner, but they can "
            "save 2% by paying early. Customers who ordered late in the day may "
            "need help to change their routine.")
    para(d, "Example: a co-op that orders at 4:00 pm on Sunday will now receive "
            "the order on Tuesday, not Monday. If it orders before 2:00 pm, "
            "nothing changes.")
    h2(d, "6. What staff should do")
    bullets(d, ["Tell every wholesale customer about the changes before 15 October 2026.",
                "Send the customer letter template (from Marketing) by email.",
                "Help regular late-day customers move to a morning order routine.",
                "Log any complaint or special request in the CRM with the tag “Nov-terms”.",
                "Do not promise exceptions. Send special requests to your manager."],
            style="List Number")
    h2(d, "7. Questions")
    para(d, "Sales questions: Noura Al-Ajmi, Head of Sales (sales@%s). Payment "
            "and invoice questions: Finance team (finance@%s). Thank you for "
            "helping our customers through this change." % (DOMAIN, DOMAIN))
    d.save(out("tamra-supplier-notice.docx"))


# ===========================================================================
# 13. Leadership meeting notes (raw)
# ===========================================================================
def build_meeting_notes():
    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    fp = d.sections[0].footer.paragraphs[0]
    fp.text = DISCLAIMER
    for r in fp.runs:
        r.font.size = Pt(8)
        r.font.italic = True
    p = d.add_paragraph()
    p.add_run("leadership mtg — Sun 27 Sep 2026 — 9:05am-10:50am, "
              "board rm Shuwaikh").bold = True
    L = [
        "present: Dana (chair), Faisal, Noura, Sara, Omar, Yousef. notes: Hind (CEO office). Omar joined late ~9:20",
        "",
        "1) H1 / Aug numbers",
        "- Faisal: delivery app rev up a lot (92k Jan -> 148k Aug) BUT op costs up faster since May. driver pay + platform fees. Aug margin approx -3%, first negative month",
        "- Dana: so we grow and lose money on each order?? need fix before Q4",
        "- wholesale July dip ~21% vs June. Yousef says hotels quiet in summer + 2 hotels paused orders (renovation). Aug partly back",
        "- Noura: top 5 hotel clients = ~31% of wholesale. risky. Dana agrees, wants plan",
        "- cafes fine, growing ~2%/mo, sat 4.6ish",
        "",
        "2) POS upgrade",
        "- Omar: Fahaheel POS tickets spiked in Aug after firmware 4.2. terminals froze, card payments failed. rolled back some, still issues",
        "- Omar asks for full upgrade all 6 cafes + integration w/ delivery app. budget 120k (in board pack)",
        "- pilot Salmiya first. Yousef ok but NOT during cold brew launch wk (18 Oct)",
        "- DECISION: go ahead w/ pilot at Salmiya, start Sun 25 Oct. full rollout Nov-Dec subject to board approval",
        "- Yousef: no more auto firmware updates until IT tests at 1 cafe. Omar ok",
        "",
        "3) Delivery app fees",
        "- Faisal: 2 options for new fee model: (a) customer delivery fee 0.500 KWD on orders under 5 KWD, (b) renegotiate partner commission from 22% -> 17%. prob both",
        "- Sara worried (a) hurts app growth / new users",
        "- DECISION: pause all new delivery promo codes until fee model agreed. existing promos run to end Oct",
        "- open Q: who leads partner negotiation — Faisal or Yousef? not agreed",
        "",
        "4) Cold brew launch (18 Oct)",
        "- Sara: budget 18k (Insta 4.5k, TikTok 4k, influencers 3k, Snap 2.5k, in-cafe 2.5k, email 0.5k, contingency 1k)",
        "- Faisal: too much given delivery margin. wants 12k, drop influencers + cut TikTok. \"we can do in-cafe + Insta only\"",
        "- Sara disagrees strongly — TikTok + creators = how we reach 22-35. without them target 25k bottles not realistic",
        "- Dana: not deciding today. both bring options to board",
        "- open Q: final cold brew budget 18k or 12k?",
        "",
        "5) Sales / Darwaza Hotels",
        "- Noura: Darwaza Hotels (4 hotels) ~70k/yr, decision by 30 Nov, start Jan. current supplier late deliveries",
        "- DECISION: Noura leads Darwaza proposal, 3 price tiers ok (Classic / Premium / Signature)",
        "- Khalid lost a lot of deals on price this yr — Noura to look at discount rules. (not an action yet?)",
        "",
        "6) AOB",
        "- Faisal: POS 102k already in Q4 capex so cash ok",
        "- Jahra cafe weekend hours — keep late opening Thu/Fri? no data. open Q",
        "- wholesale July dip — one-off (summer) or trend? Yousef thinks one-off. Dana wants proof. open Q",
        "",
        "ACTIONS",
        "- Faisal — delivery fee model options (a)+(b) w/ numbers for board pack — by Thu 1 Oct",
        "- Omar — POS pilot plan + vendor contract for Salmiya — next wk",
        "- Noura — Darwaza proposal draft to Dana — 15 Oct",
        "- Sara — cold brew budget, 2 options (18k / 12k) + impact on targets — before board",
        "- Yousef — wholesale July dip analysis + call top 5 hotel clients — end of month-ish",
        "",
        "next mtg Sun 4 Oct same time. board mtg date tbc (mid Oct?)",
    ]
    for line in L:
        d.add_paragraph(line)
    d.save(out("tamra-leadership-meeting-notes.docx"))
    FACTS["meeting_notes"] = {
        "file": "tamra-leadership-meeting-notes.docx",
        "meeting": "Sunday 27 Sep 2026, 9:05-10:50, leadership team",
        "planted_contradiction": {
            "topic": "POS upgrade budget",
            "value_1": "KWD 120,000 (section 2: 'budget 120k (in board pack)')",
            "value_2": "KWD 102,000 (section 6 AOB: 'POS 102k already in Q4 capex')",
            "correct_value_per_other_files": 120000,
        },
        "decisions": [
            "Go ahead with the POS pilot at Salmiya, starting Sunday 25 Oct 2026 (full rollout Nov-Dec subject to board approval)",
            "Pause all new delivery promo codes until the new fee model is agreed (existing promos run to end of October)",
            "Noura leads the Darwaza Hotels proposal; the 3 price tiers (Classic / Premium / Signature) are approved",
        ],
        "open_questions": [
            "Who leads the delivery-partner negotiation: Faisal or Yousef?",
            "Final cold brew budget: KWD 18,000 or KWD 12,000?",
            "Keep late weekend opening hours at the Jahra cafe?",
            "Is the wholesale July dip a one-off (summer) or a trend?",
        ],
        "actions": [
            {"owner": "Faisal Al-Rashed (CFO)", "action": "Delivery fee model options (a)+(b) with numbers for the board pack", "due": "Thu 1 Oct 2026"},
            {"owner": "Omar Al-Enezi (Head of IT)", "action": "POS pilot plan and vendor contract for Salmiya", "due": "'next wk' (vague; week of 4 Oct 2026)"},
            {"owner": "Noura Al-Ajmi (Head of Sales)", "action": "Darwaza Hotels proposal draft to Dana", "due": "15 Oct 2026"},
            {"owner": "Sara Al-Hajri (Head of Marketing)", "action": "Cold brew budget, 2 options (18k / 12k) with impact on targets", "due": "'before board' (vague; board date not set)"},
            {"owner": "Yousef Al-Shammari (COO)", "action": "Wholesale July dip analysis and calls to the top 5 hotel clients", "due": "'end of month-ish' (vague; ~30 Sep 2026)"},
        ],
        "disagreement": "CFO Faisal wants the cold brew budget cut to KWD 12,000 (drop influencers, cut TikTok); Head of Marketing Sara says KWD 18,000 is needed to reach 22-35-year-olds and hit 25,000 bottles. Dana deferred to the board.",
    }


# ===========================================================================
# Verification: re-open every file and print a summary
# ===========================================================================
def verify():
    print("\nFile summary")
    print("-" * 72)
    summary = []
    for name in WRITTEN:
        path = os.path.join(SITE, name)
        size = os.path.getsize(path)
        if name.endswith(".xlsx"):
            wb = load_workbook(path)
            ws = wb.worksheets[0]
            assert "About" in wb.sheetnames
            assert any(DISCLAIMER in str(c.value) for c in wb["About"]["A"] if c.value)
            tables = list(ws.tables.keys())
            info = "sheet '%s', table %s, %d data rows x %d cols" % (
                ws.title, tables, ws.max_row - 1, ws.max_column)
        elif name.endswith(".docx"):
            dd = docx.Document(path)
            text = " ".join(p.text for p in dd.paragraphs)
            for t in dd.tables:
                for r in t.rows:
                    text += " " + " ".join(c.text for c in r.cells)
            words = len(text.split())
            assert DISCLAIMER in dd.sections[0].footer.paragraphs[0].text
            heads = sum(1 for p in dd.paragraphs if p.style.name.startswith("Heading"))
            info = "%d words, %d headings, %d tables, ~%.1f pages" % (
                words, heads, len(dd.tables), words / 420.0 + 0.15 * len(dd.tables))
            if name == "tamra-company-profile.docx":
                low = text.lower()
                for bad in ("profit", "margin", "ebitda", "net income"):
                    assert bad not in low, bad
        elif name.endswith(".pptx"):
            pr = Presentation(path)
            assert DISCLAIMER in pr.slides[0].notes_slide.notes_text_frame.text
            titles = [s.shapes.title.text for s in pr.slides]
            info = "%d slides: %s" % (len(pr.slides), " | ".join(titles))
        else:
            info = ""
        summary.append((name, size, info))
        print("%-42s %7.1f KB  %s" % (name, size / 1024.0, info))
    return summary


def main():
    assert os.path.isdir(SITE), SITE
    # Day 2 Excel first (docs reuse the numbers)
    kpi_rows = build_kpis()
    by, h1r, h1t = kpi_facts(kpi_rows)
    deals = build_pipeline()
    pipeline_facts(deals)
    camp_rows = build_campaigns()
    campaign_facts(camp_rows)
    tickets = build_tickets()
    fah_aug = ticket_facts(tickets)
    # Word / PowerPoint
    build_company_profile()
    build_h1_review(by, h1r, h1t)
    build_account_brief()
    build_cold_brew_brief()
    build_pos_plan(fah_aug)
    build_messy_deck()
    build_house_letter()
    build_supplier_notice()
    build_meeting_notes()

    order = ["tamra-company-profile.docx", "tamra-monthly-kpis.xlsx",
             "tamra-sales-pipeline.xlsx", "tamra-campaign-results.xlsx",
             "tamra-it-tickets.xlsx", "tamra-h1-2026-business-review.docx",
             "tamra-account-brief-darwaza-hotels.docx",
             "tamra-campaign-brief-cold-brew.docx",
             "tamra-pos-upgrade-project-plan.docx", "tamra-messy-deck.pptx",
             "tamra-house-style-letter.docx", "tamra-supplier-notice.docx",
             "tamra-leadership-meeting-notes.docx"]
    assert sorted(order) == sorted(WRITTEN)
    WRITTEN[:] = order
    FACTS["_about"] = {
        "company": "Tamra Foods Co. (fictional)",
        "disclaimer": DISCLAIMER,
        "generator": "build/assets-src/make_assets.py (random.seed(2026))",
        "files": order,
    }
    with open(FACTS_PATH, "w", encoding="utf-8") as f:
        json.dump(FACTS, f, indent=2, ensure_ascii=False, default=str)
    verify()
    print("\nfacts.json written to", FACTS_PATH)


if __name__ == "__main__":
    main()
